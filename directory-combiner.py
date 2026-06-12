#!/usr/bin/env python3
"""Combine all XLSX files in this script's directory.

Workflow:
1. Look only in the folder where this script is located (no subfolders).
2. For every input .xlsx file, copy the value from cell A2 down through the
   worksheet's last row (A3, A4, etc.), then save that file.
3. Combine all input workbooks into one workbook named ``combined.xlsx``.
   The first file keeps its header row; every later file starts from row 2 so
   duplicate headers are not included.

This script is intended to be run directly from Thonny. It requires openpyxl.
In Thonny, install it with: Tools -> Manage packages -> search "openpyxl".
"""

from pathlib import Path
import sys
from typing import List

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("This script needs the 'openpyxl' package to work with .xlsx files.")
    print('In Thonny: Tools -> Manage packages -> search for "openpyxl" -> Install.')
    input("Press Enter to close...")
    sys.exit(1)


OUTPUT_FILE_NAME = "combined.xlsx"
TEMP_FILE_PREFIX = "~$"


def script_directory() -> Path:
    """Return the directory that contains this script."""
    try:
        return Path(__file__).resolve().parent
    except NameError:
        # Fallback for unusual interactive runners.
        return Path.cwd().resolve()


def find_input_files(folder: Path) -> List[Path]:
    """Find .xlsx files directly inside *folder*, excluding temporary/output files."""
    return sorted(
        file_path
        for file_path in folder.glob("*.xlsx")
        if file_path.is_file()
        and not file_path.name.startswith(TEMP_FILE_PREFIX)
        and file_path.name.lower() != OUTPUT_FILE_NAME.lower()
    )


def last_data_row(worksheet) -> int:
    """Return the last row that contains any real cell value."""
    for row_number in range(worksheet.max_row, 0, -1):
        for cell in worksheet[row_number]:
            if cell.value not in (None, ""):
                return row_number
    return 0


def fill_column_a_from_a2(file_path: Path) -> None:
    """Copy the active sheet's A2 value down to the active sheet's last data row."""
    workbook = load_workbook(file_path)
    worksheet = workbook.active

    value_to_copy = worksheet["A2"].value
    final_entry_row = last_data_row(worksheet)

    if final_entry_row >= 3:
        for row_number in range(3, final_entry_row + 1):
            worksheet.cell(row=row_number, column=1).value = value_to_copy

    workbook.save(file_path)
    workbook.close()


def append_workbook_rows(source_file: Path, output_sheet, include_header: bool) -> int:
    """Append rows from a source workbook into the output sheet.

    Returns the number of rows appended.
    """
    workbook = load_workbook(source_file, data_only=False)
    worksheet = workbook.active
    start_row = 1 if include_header else 2
    rows_added = 0

    for row in worksheet.iter_rows(min_row=start_row, values_only=True):
        output_sheet.append(row)
        rows_added += 1

    workbook.close()
    return rows_added


def combine_files(input_files: List[Path], output_file: Path) -> int:
    """Combine all input files into one output workbook and return rows added."""
    combined_workbook = Workbook()
    combined_sheet = combined_workbook.active
    combined_sheet.title = "Combined"

    total_rows_added = 0
    for index, input_file in enumerate(input_files):
        total_rows_added += append_workbook_rows(
            input_file,
            combined_sheet,
            include_header=(index == 0),
        )

    combined_workbook.save(output_file)
    combined_workbook.close()
    return total_rows_added


def main() -> int:
    folder = script_directory()
    output_file = folder / OUTPUT_FILE_NAME
    input_files = find_input_files(folder)

    if not input_files:
        print(f"No input .xlsx files found in: {folder}")
        print(f"Put this script in the same folder as your .xlsx files and run it again.")
        input("Press Enter to close...")
        return 1

    print(f"Working folder: {folder}")
    print(f"Found {len(input_files)} .xlsx file(s) to process:")
    for input_file in input_files:
        print(f"- {input_file.name}")

    print("\nStep 1: Copying A2 down column A in each file...")
    for input_file in input_files:
        fill_column_a_from_a2(input_file)
        print(f"Updated: {input_file.name}")

    print(f"\nStep 2: Combining files into {OUTPUT_FILE_NAME}...")
    rows_added = combine_files(input_files, output_file)

    print("\nDone!")
    print(f"Created: {output_file}")
    print(f"Rows written to combined file: {rows_added}")
    input("Press Enter to close...")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
